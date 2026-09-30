"""Original short-term inertial estimates; never claim absolute positioning."""
import csv
from io import BytesIO, StringIO
import math

from .waveform_export import _matrix, _text

COLUMNS = ("elapsed_s", "measurement_time_s", "host_receive_time_s", "segment",
           "position_x_m", "position_y_m", "position_z_m",
           "velocity_x_m_s", "velocity_y_m_s", "velocity_z_m_s",
           "acceleration_x_m_s2", "acceleration_y_m_s2", "acceleration_z_m_s2",
           "distance_m", "stationary", "qw", "qx", "qy", "qz", "orientation_kind")
ORIENTATION = {"quaternion": 1, "euler_zyx": 2}


def finite(value):
    return type(value) in (int, float) and math.isfinite(value)


def vector(value, count=3):
    return isinstance(value, list) and len(value) == count and all(map(finite, value))


def export(data):
    if data.get("estimate") is not True:
        raise ValueError("轨迹导出必须标明 estimate=true")
    if set(data) - {"format", "source", "generation", "gravity_reference", "timing", "points", "estimate"}:
        raise ValueError("包含未知轨迹导出参数")
    kind = data.get("format")
    if kind not in {"mat", "xlsx", "csv"}:
        raise ValueError("轨迹导出格式必须为 mat、xlsx 或 csv")
    if data.get("source") not in {"live", "demo", "playback", "none"} or data.get("timing") not in {"host_receive_time", "recorded_time"}:
        raise ValueError("无效的轨迹数据来源或时基")
    if type(data.get("generation")) is not int or not vector(data.get("gravity_reference")):
        raise ValueError("无效的来源代次或重力参考")
    points = data.get("points")
    if not isinstance(points, list) or not 1 <= len(points) <= 12000:
        raise ValueError("单次轨迹导出需要 1–12000 个估算点")
    matrix = []
    previous = -math.inf
    for point in points:
        if not isinstance(point, dict) or set(point) - {"time", "host_time", "elapsed", "segment", "position", "velocity", "acceleration", "distance", "stationary", "orientation_kind", "quaternion"}:
            raise ValueError("无效的轨迹点")
        if any(not finite(point.get(key)) for key in ("time", "host_time", "elapsed", "distance")):
            raise ValueError("轨迹时间和距离必须为有限值")
        if point["elapsed"] < previous or point["elapsed"] < 0 or point["distance"] < 0:
            raise ValueError("轨迹时间顺序或距离无效")
        previous = point["elapsed"]
        if type(point.get("segment")) is not int or point["segment"] < 0 or type(point.get("stationary")) is not bool:
            raise ValueError("无效的轨迹段或静止标记")
        if not all(vector(point.get(key)) for key in ("position", "velocity", "acceleration")) or not vector(point.get("quaternion"), 4):
            raise ValueError("无效的轨迹向量或姿态")
        if point.get("orientation_kind") not in ORIENTATION:
            raise ValueError("未知的轨迹姿态来源")
        matrix.append([point["elapsed"], point["time"], point["host_time"], point["segment"],
                       *point["position"], *point["velocity"], *point["acceleration"],
                       point["distance"], int(point["stationary"]), *point["quaternion"], ORIENTATION[point["orientation_kind"]]])
    metadata = {"estimate": "Short-term inertial estimate; drift accumulates; not absolute position",
                "source": data["source"], "generation": data["generation"], "timing": data["timing"],
                "coordinate_frame": "right-handed world XYZ; Z-up; origin at tracking start",
                "gravity_reference": ",".join(map(str, data["gravity_reference"])),
                "orientation_kind": "1=quaternion,2=euler_zyx"}
    if kind == "csv":
        out = StringIO(newline="");writer = csv.writer(out)
        writer.writerow((*COLUMNS, "source", "timing", "estimate", "gravity_reference_x", "gravity_reference_y", "gravity_reference_z"))
        for row in matrix:writer.writerow((*row, data["source"], data["timing"], True, *data["gravity_reference"]))
        return b"\xef\xbb\xbf" + out.getvalue().encode(), "text/csv; charset=utf-8"
    if kind == "mat":
        out = BytesIO(b"MATLAB 5.0 MAT-file, DM IMU Workbench inertial estimate".ljust(116, b" ") + b"\0" * 8 + b"\0\x01IM")
        out.seek(0, 2);out.write(_matrix("trajectory", matrix, len(COLUMNS)))
        out.write(_text("trajectory_columns", ",".join(COLUMNS)))
        out.write(_matrix("gravity_reference", [data["gravity_reference"]], 3))
        for key, value in metadata.items():out.write(_text(key + "_text" if key == "gravity_reference" else key, str(value)))
        return out.getvalue(), "application/x-matlab-data"
    from openpyxl import Workbook
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("metadata");sheet.append(("key", "value"))
    for key, value in metadata.items():sheet.append((key, value))
    sheet = workbook.create_sheet("trajectory");sheet.freeze_panes = "E2";sheet.append(COLUMNS)
    for row in matrix:sheet.append(row)
    out = BytesIO();workbook.save(out)
    return out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
