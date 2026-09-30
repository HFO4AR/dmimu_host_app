"""Export one-sided, gain-corrected FFT amplitudes with their analysis settings."""
import csv
from io import BytesIO, StringIO
import math

from .waveform_export import CHANNELS, _matrix, _text


def export(data):
    allowed = {"format", "channel", "source", "sample_rate_hz", "window", "samples", "start_time_unix_s", "end_time_unix_s", "frequency_hz", "amplitude"}
    if not isinstance(data, dict) or set(data) != allowed:
        raise ValueError("频谱导出参数不完整或包含未知字段")
    kind, channel = data["format"], data["channel"]
    if kind not in {"mat", "xlsx", "csv"} or channel not in CHANNELS:
        raise ValueError("无效的频谱格式或通道")
    n, fs = data["samples"], data["sample_rate_hz"]
    if type(n) is not int or n < 16 or n > 65536 or n & (n - 1):
        raise ValueError("FFT 点数必须为 16–65536 的 2 的幂")
    if type(fs) not in (int, float) or not math.isfinite(fs) or not 1 <= fs <= 100000:
        raise ValueError("无效采样率")
    if data["window"] not in {"hann", "hamming", "rectangular"} or data["source"] not in {"live", "demo", "playback", "unknown", "none"}:
        raise ValueError("无效窗函数或数据来源")
    for name in ("start_time_unix_s", "end_time_unix_s"):
        if type(data[name]) not in (int, float) or not math.isfinite(data[name]):
            raise ValueError("无效频谱时间范围")
    if data["end_time_unix_s"] < data["start_time_unix_s"]:
        raise ValueError("无效频谱时间顺序")
    keys, unit = CHANNELS[channel]
    frequencies, amplitudes = data["frequency_hz"], data["amplitude"]
    if not isinstance(frequencies, list) or len(frequencies) != n // 2 + 1 or not isinstance(amplitudes, list) or len(amplitudes) != len(keys):
        raise ValueError("频谱数组的维度不匹配")
    if any(not isinstance(axis, list) or len(axis) != len(frequencies) for axis in amplitudes):
        raise ValueError("频谱轴的点数不匹配")
    for i, value in enumerate(frequencies):
        if type(value) not in (int, float) or not math.isfinite(value) or abs(value - i * fs / n) > max(1e-9, fs * 1e-10):
            raise ValueError("频率轴必须为 0 至 Fs/2 的 FFT 原始频点")
    if any(type(v) not in (int, float) or not math.isfinite(v) or v < 0 for axis in amplitudes for v in axis):
        raise ValueError("频谱幅值必须为非负有限数值")
    columns = ("frequency_hz", *("amplitude_" + key for key in keys))
    rows = list(zip(frequencies, *amplitudes))
    metadata = {key: data[key] for key in allowed - {"format", "frequency_hz", "amplitude"}}
    metadata.update(unit=unit, amplitude_definition="one-sided peak amplitude; mean removed; coherent window gain corrected; not PSD")
    if kind == "mat":
        out = BytesIO(b"MATLAB 5.0 MAT-file, DM IMU Workbench FFT spectrum".ljust(116, b" ") + b"\0" * 8 + b"\0\x01IM")
        out.seek(0, 2)
        out.write(_matrix("spectrum", rows, len(columns)))
        out.write(_text("spectrum_columns", ",".join(columns)))
        for key, value in sorted(metadata.items()):
            out.write(_text(key, value) if isinstance(value, str) else _matrix(key, [[value]], 1))
        return out.getvalue(), "application/x-matlab-data"
    if kind == "csv":
        out = StringIO(newline=""); writer = csv.writer(out)
        # A rectangular table remains directly usable by readtable/pandas/Excel.
        meta = sorted(metadata)
        writer.writerow((*columns, *meta))
        for row in rows:
            writer.writerow((*row, *(metadata[key] for key in meta)))
        return b"\xef\xbb\xbf" + out.getvalue().encode(), "text/csv; charset=utf-8"
    from openpyxl import Workbook
    workbook = Workbook(write_only=True)
    info = workbook.create_sheet("metadata")
    info.append(("key", "value"))
    for item in sorted(metadata.items()):
        info.append(item)
    sheet = workbook.create_sheet("spectrum"); sheet.freeze_panes = "B2"
    sheet.append(columns)
    for row in rows:
        sheet.append(row)
    out = BytesIO(); workbook.save(out)
    return out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
