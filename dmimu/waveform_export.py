"""Export original waveform samples, without interpolation or decimation."""
import csv
from io import BytesIO, StringIO
import math
import struct

CHANNELS = {
    "acceleration": (("x", "y", "z"), "m/s^2"),
    "angular_velocity": (("x", "y", "z"), "rad/s"),
    "euler": (("roll", "pitch", "yaw"), "deg"),
    "quaternion": (("w", "x", "y", "z"), "dimensionless"),
}
MAX_SAMPLES = 400000


def validate(data):
    if set(data) - {"format", "channels", "source"}:
        raise ValueError("包含未知导出参数")
    if data.get("format") not in ("mat", "xlsx", "csv"):
        raise ValueError("导出格式必须为 mat、xlsx 或 csv")
    channels = data.get("channels")
    if not isinstance(channels, dict) or set(channels) - CHANNELS.keys():
        raise ValueError("无效的波形通道")
    source = data.get("source", "unknown")
    if source not in ("live", "demo", "playback", "none", "unknown"):
        raise ValueError("无效的数据来源")
    count = 0
    start = math.inf
    for name, rows in channels.items():
        if not isinstance(rows, list):
            raise ValueError("采样数据必须为数组")
        count += len(rows)
        if count > MAX_SAMPLES:
            raise ValueError("单次波形导出最多 40 万个原始采样点")
        previous = -math.inf
        for row in rows:
            if not isinstance(row, list) or len(row) != len(CHANNELS[name][0]) + 1:
                raise ValueError("采样行的列数不匹配")
            if any(type(v) not in (int, float) or not math.isfinite(v) for v in row):
                raise ValueError("采样值必须是有限数值")
            if row[0] < previous:
                raise ValueError("采样时间必须按顺序排列")
            previous = row[0]
            start = min(start, previous)
    if not count:
        raise ValueError("当前区间没有可导出的原始采样点")
    return channels, source, start


def _tag(kind, payload):
    return struct.pack("<II", kind, len(payload)) + payload + b"\0" * (-len(payload) % 8)


def _matrix(name, rows, columns):
    # Level 5 miMATRIX, mxDOUBLE_CLASS, column-major IEEE754 doubles.
    body = _tag(6, struct.pack("<II", 6, 0))
    body += _tag(5, struct.pack("<ii", len(rows), columns))
    body += _tag(1, name.encode("ascii"))
    values = BytesIO()
    for col in range(columns):
        for row in rows:
            values.write(struct.pack("<d", row[col]))
    return _tag(14, body + _tag(9, values.getvalue()))


def _text(name, value):
    raw = value.encode("utf-16-le")
    body = _tag(6, struct.pack("<II", 4, 0))
    body += _tag(5, struct.pack("<ii", 1, len(raw) // 2))
    body += _tag(1, name.encode("ascii"))
    return _tag(14, body + _tag(4, raw))


def export(data):
    channels, source, origin = validate(data)
    kind = data["format"]
    if kind == "mat":
        description = b"MATLAB 5.0 MAT-file, DM IMU Workbench original waveform samples"
        out = BytesIO(description.ljust(116, b" ") + b"\0" * 8 + b"\0\x01IM")
        out.seek(0, 2)
        out.write(_matrix("time_origin_unix_s", [[origin]], 1))
        out.write(_text("source", source))
        out.write(_text("time_note", "time_s = host_time_unix_s - time_origin_unix_s; host receive time, not device clock"))
        for name, rows in channels.items():
            axes, unit = CHANNELS[name]
            matrix = [[row[0] - origin, *row] for row in rows]
            out.write(_matrix(name, matrix, len(axes) + 2))
            out.write(_text(name + "_columns", ",".join(("time_s", "host_time_unix_s", *axes))))
            out.write(_text(name + "_unit", unit))
        return out.getvalue(), "application/x-matlab-data"
    if kind == "csv":
        out = StringIO(newline="")
        writer = csv.writer(out)
        axes = ("x", "y", "z", "w", "roll", "pitch", "yaw")
        writer.writerow(("time_s", "host_time_unix_s", "source", "channel", "unit", *axes))
        for name, rows in channels.items():
            keys, unit = CHANNELS[name]
            for row in rows:
                values = dict(zip(keys, row[1:]))
                writer.writerow((row[0] - origin, row[0], source, name, unit, *(values.get(key, "") for key in axes)))
        return b"\xef\xbb\xbf" + out.getvalue().encode("utf-8"), "text/csv; charset=utf-8"
    from openpyxl import Workbook
    workbook = Workbook(write_only=True)
    info = workbook.create_sheet("metadata")
    for row in (("key", "value"), ("source", source), ("time_origin_unix_s", origin),
                ("sampling", "Original samples; no decimation or interpolation"),
                ("clock", "Host receive timestamps; not device clock"),
                ("angular_velocity_unit", "rad/s")):
        info.append(row)
    for name, rows in channels.items():
        keys, unit = CHANNELS[name]
        sheet = workbook.create_sheet(name)
        sheet.freeze_panes = "C2"
        sheet.append(("time_s", "host_time_unix_s", *(key + " [" + unit + "]" for key in keys)))
        for row in rows:
            sheet.append((row[0] - origin, *row))
    out = BytesIO()
    workbook.save(out)
    return out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
